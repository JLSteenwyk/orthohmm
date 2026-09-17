"""Freeze the complete experimental WGD cohort without reading tool predictions."""

import argparse
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from benchmark_tools.snapshot_orthohmm_input_order import record

HEADERS = ("DM strainID", "SM1 strainID", "SM2 strainID", "ORF1", "ORF2", "Gene1", "Gene2",
           "TGI", "DGI-1", "DGI-2", "TGI frac", "DGI-1 frac", "DGI-2 frac",
           "Meets degree threshold", "Trigenic interaction fraction class")


def parse_rows(rows):
    result, used = [], set()
    for row_number, row in enumerate(rows, 3):
        if len(row) != 15:
            raise ValueError("Unexpected experimental table width")
        genes = list(row[3:5])
        if (any(not isinstance(g, str) or not re.fullmatch(r"Y[A-P][LR]\d{3}[CW](?:-[A-Z])?", g) for g in genes)
                or genes[0] == genes[1] or used.intersection(genes)):
            raise ValueError("Invalid, duplicate or reused experimental ORF")
        used.update(genes)
        counts = row[7:10]
        if any(type(n) is not int or n < 0 for n in counts):
            raise ValueError("Invalid interaction counts")
        eligible = max(counts) >= 6
        if row[13] not in (0, 1) or bool(row[13]) != eligible:
            raise ValueError("Source degree flag disagrees with its stated criterion")
        if eligible:
            total = sum(counts)
            fractions = list(row[10:13])
            if any(not isinstance(f, (int, float)) or not math.isfinite(f)
                   or not math.isclose(f, n / total, rel_tol=1e-10, abs_tol=1e-12)
                   for f, n in zip(fractions, counts)):
                raise ValueError("Source fractions do not reproduce from counts")
            fraction = fractions[0]
            category = "High" if fraction > .4 else "Low" if fraction < .4 else "Boundary"
            if row[14] != category or category == "Boundary":
                raise ValueError("Unresolved or inconsistent source class boundary")
        else:
            if any(value != "NaN" for value in row[10:13]) or row[14] is not None:
                raise ValueError("Sparse measurements must remain explicitly unclassified")
            fraction, category = None, "Sparse"
        pair = sorted(genes)
        names = dict(zip(genes, row[5:7]))
        result.append({"orf_pair": pair, "gene_names": [names[g] for g in pair],
                       "experimental_class": category, "trigenic_fraction": fraction,
                       "trigenic_count": counts[0], "digenic_counts_by_orf": dict(zip(genes, counts[1:])),
                       "source_row": row_number,
                       "example_rank_sha256": hashlib.sha256("\t".join(pair).encode("ascii")).hexdigest()})
    return sorted(result, key=lambda row: row["orf_pair"])


def selected_examples(rows):
    selected = []
    for category in ("High", "Low", "Sparse"):
        candidates = sorted((row for row in rows if row["experimental_class"] == category),
                            key=lambda row: row["example_rank_sha256"])
        if len(candidates) < 2:
            raise ValueError("Too few source pairs for frozen example selection")
        selected.extend({"orf_pair": row["orf_pair"], "experimental_class": category,
                         "rank_sha256": row["example_rank_sha256"]} for row in candidates[:2])
    return selected


def freeze(root, protocol):
    metadata_path, table_path = root / "zenodo_3975054.json", root / "TableS7.xlsx"
    metadata_record, table_record = record(metadata_path), record(table_path)
    if metadata_record["sha256"] != "a113657e4a5b914fb2c06ddaa3fc35c941b995c7e3fc1c3db00d7d6a4c3222e8":
        raise ValueError("Changed captured deposit metadata")
    if table_record["sha256"] != "bf502168d9c87956edf6c764650e45a8c60a8605e08ea93fd7ef57a49313f7fe":
        raise ValueError("Changed prespecified experimental table")
    metadata = json.loads(metadata_path.read_text())
    entries = [row for row in metadata["files"] if row["key"] == "TableS7.xlsx"]
    if len(entries) != 1 or entries[0]["size"] != table_record["bytes"]:
        raise ValueError("Deposit table inventory mismatch")
    if entries[0]["checksum"] != "md5:" + hashlib.md5(table_path.read_bytes()).hexdigest():
        raise ValueError("Deposited checksum differs")
    if metadata["id"] != 3975054 or metadata["metadata"]["license"]["id"] != "cc-zero":
        raise ValueError("Unexpected source identity/license")
    with table_path.open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        if workbook.sheetnames != ["Trigenic fraction list"]:
            raise ValueError("Unexpected workbook sheets")
        sheet = workbook.active
        headers = tuple(str(v).strip() for v in next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        if headers != HEADERS:
            raise ValueError("Experimental table schema changed")
        rows = parse_rows(list(sheet.iter_rows(min_row=3, values_only=True)))
        workbook.close()
    counts = dict(Counter(row["experimental_class"] for row in rows))
    if len(rows) != 240 or counts != {"High": 47, "Low": 114, "Sparse": 79}:
        raise ValueError("Complete source cohort/class inventory differs")
    return {"status": "experimental_wgd_cohort_frozen_no_predictions_read", "pairs": rows,
            "pair_count": len(rows), "unique_orfs": len({g for row in rows for g in row["orf_pair"]}),
            "class_counts": counts, "prespecified_examples": selected_examples(rows),
            "metadata": metadata_record, "table": table_record, "protocol": record(protocol),
            "source": record(__file__), "openpyxl_version": openpyxl.__version__,
            "download_url": entries[0]["links"]["self"], "deposit_license": "CC0-1.0",
            "limitations": ["Functional interaction measurements define strata, not orthology truth.",
                            "Sparse source fractions remain missing, not zero.",
                            "No prediction files read; reference/input mapping and biological outcomes remain unevaluated."]}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    for flag in ("root", "protocol", "output"):
        parser.add_argument("--" + flag, type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(args.output)
    result = freeze(args.root.resolve(), args.protocol.resolve())
    with args.output.open("x") as handle:
        json.dump(result, handle, indent=2, sort_keys=True)
        handle.write("\n")
